import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # Set slide dimensions to 16:9 widescreen (13.333 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Palette
    C_NAVY_DARK = RGBColor(10, 37, 64)       # Primary Dark Navy
    C_NAVY_DEEP = RGBColor(15, 23, 42)       # Midnight Slate
    C_GOLD = RGBColor(244, 180, 26)          # Brand Amber / Gold
    C_GOLD_LIGHT = RGBColor(254, 243, 199)   # Gold tint
    C_BLUE_LIGHT = RGBColor(238, 242, 255)   # Light blue tint
    C_WHITE = RGBColor(255, 255, 255)        # Pure White
    C_GRAY_LIGHT = RGBColor(248, 250, 252)   # Card bg light
    C_GRAY_CARD = RGBColor(241, 245, 249)    # Border/card bg
    C_TEXT_DARK = RGBColor(30, 41, 59)       # Dark text
    C_TEXT_MUTED = RGBColor(100, 116, 139)   # Muted gray text
    C_ACCENT_GREEN = RGBColor(16, 185, 129)  # Green badge
    C_BORDER = RGBColor(226, 232, 240)

    def add_header(slide, title_text, category_text="TIRUMALA DIGITAL PLATFORM"):
        # Top accent bar
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = C_GOLD
        bar.line.fill.background()

        # Category pill/label
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.name = 'Calibri'
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = C_GOLD

        # Slide Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.5), Inches(0.7))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = 'Calibri'
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = C_NAVY_DARK

    def add_card(slide, left, top, width, height, bg_color=C_WHITE, border_color=C_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(1)
        else:
            card.line.fill.background()
        return card

    # ==========================================
    # SLIDE 1: Title Slide (Dark Theme)
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = C_NAVY_DARK
    bg1.line.fill.background()

    # Decorative top bar
    bar1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.2))
    bar1.fill.solid()
    bar1.fill.fore_color.rgb = C_GOLD
    bar1.line.fill.background()

    # Content container
    t_box = s1.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(11.3), Inches(4.5))
    tf1 = t_box.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "OFFICIAL DIGITAL PLATFORM SHOWCASE"
    p0.font.name = 'Calibri'
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = C_GOLD
    p0.space_after = Pt(14)

    p1 = tf1.add_paragraph()
    p1.text = "Tirumala Educational Institutions"
    p1.font.name = 'Calibri'
    p1.font.size = Pt(40)
    p1.font.bold = True
    p1.font.color.rgb = C_WHITE
    p1.space_after = Pt(8)

    p2 = tf1.add_paragraph()
    p2.text = "Next-Generation Web Experience & Multi-Campus Digital Infrastructure"
    p2.font.name = 'Calibri'
    p2.font.size = Pt(20)
    p2.font.color.rgb = RGBColor(203, 213, 225)
    p2.space_after = Pt(28)

    p3 = tf1.add_paragraph()
    p3.text = "Bhimavaram  •  Rajahmundry  •  Kakinada  •  Vijayawada  •  Guntur"
    p3.font.name = 'Calibri'
    p3.font.size = Pt(14)
    p3.font.bold = True
    p3.font.color.rgb = C_GOLD
    p3.space_after = Pt(24)

    # Feature badges card at bottom of slide 1
    badge_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(5.3), Inches(11.333), Inches(1.4))
    badge_card.fill.solid()
    badge_card.fill.fore_color.rgb = RGBColor(19, 47, 76)
    badge_card.line.color.rgb = RGBColor(30, 64, 100)
    badge_card.line.width = Pt(1)

    b_tf = badge_card.text_frame
    b_tf.word_wrap = True
    bp = b_tf.paragraphs[0]
    bp.text = "KEY PLATFORM CAPABILITIES AT A GLANCE"
    bp.font.name = 'Calibri'
    bp.font.size = Pt(11)
    bp.font.bold = True
    bp.font.color.rgb = C_GOLD
    bp.space_after = Pt(6)

    bp2 = b_tf.add_paragraph()
    bp2.text = "✔ Universal Multi-Data Search Engine    ✔ Instant Dark & Bright Mode    ✔ 5-Campus Direct Call Modal\n✔ Verified Ranks & Results Showcase     ✔ BIEAP Model Papers Portal      ✔ Responsive Mobile-First Architecture"
    bp2.font.name = 'Calibri'
    bp2.font.size = Pt(13)
    bp2.font.color.rgb = C_WHITE

    # ==========================================
    # SLIDE 2: Executive Vision & Platform Objectives
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Executive Vision & Strategic Objectives", "Strategic Transformation")

    cards_data_s2 = [
        ("01", "Institutional Authority", "Project Tirumala as the premier destination for Intermediate, NEET, JEE Advanced, and Degree coaching across coastal Andhra Pradesh with pristine brand aesthetics."),
        ("02", "Seamless Admissions", "Eliminate parent friction by placing direct campus phone lines, WhatsApp academic counseling, and application forms directly at the user's fingertips."),
        ("03", "Academic Transparency", "Instant online verification of state top ranks, roll numbers, BIEAP model question papers, blueprints, and examination notices."),
        ("04", "Modern Digital Standards", "Sub-second load times, instant dark/light theme switching, intelligent universal search, and flawless mobile experience on every device.")
    ]

    for i, (num, title, desc) in enumerate(cards_data_s2):
        left = Inches(0.8 + i * 2.95)
        top = Inches(1.8)
        w = Inches(2.8)
        h = Inches(4.8)

        add_card(s2, left, top, w, h, C_WHITE, C_BORDER)

        # Number badge
        nb = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(0.25), top + Inches(0.3), Inches(0.8), Inches(0.5))
        nb.fill.solid()
        nb.fill.fore_color.rgb = C_BLUE_LIGHT
        nb.line.fill.background()
        nb_tf = nb.text_frame
        nb_p = nb_tf.paragraphs[0]
        nb_p.text = num
        nb_p.alignment = PP_ALIGN.CENTER
        nb_p.font.name = 'Calibri'
        nb_p.font.size = Pt(16)
        nb_p.font.bold = True
        nb_p.font.color.rgb = C_NAVY_DARK

        # Text
        tb = s2.shapes.add_textbox(left + Inches(0.2), top + Inches(1.1), w - Inches(0.4), h - Inches(1.3))
        tf = tb.text_frame
        tf.word_wrap = True
        tp1 = tf.paragraphs[0]
        tp1.text = title
        tp1.font.name = 'Calibri'
        tp1.font.size = Pt(18)
        tp1.font.bold = True
        tp1.font.color.rgb = C_NAVY_DARK
        tp1.space_after = Pt(12)

        tp2 = tf.add_paragraph()
        tp2.text = desc
        tp2.font.name = 'Calibri'
        tp2.font.size = Pt(13)
        tp2.font.color.rgb = C_TEXT_MUTED

    # ==========================================
    # SLIDE 3: Visual Identity & Dynamic Theme System
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Modern Aesthetics: Instant Dark & Bright Modes", "User Experience Innovation")

    # Left Column: Bright Mode Showcase Card
    c_bright = add_card(s3, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8), C_WHITE, C_BORDER)
    tb_b = s3.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.0), Inches(4.3))
    tf_b = tb_b.text_frame
    tf_b.word_wrap = True

    p_b0 = tf_b.paragraphs[0]
    p_b0.text = "☀️ Pristine Bright Mode"
    p_b0.font.name = 'Calibri'
    p_b0.font.size = Pt(20)
    p_b0.font.bold = True
    p_b0.font.color.rgb = C_NAVY_DARK
    p_b0.space_after = Pt(10)

    p_b1 = tf_b.add_paragraph()
    p_b1.text = "• Tailored for Daytime & Official Reading:\n  High-contrast crisp typography, official navy `#0A2540` accents, and warm gold `#F4B41A` branding.\n\n• Zero-Flicker Architecture:\n  Inline DOM script executes before paint to prevent dark-to-light flash on reloads.\n\n• LocalStorage Memory:\n  Remembers the student or parent's preference automatically across visits.\n\n• Accessible Contrast Standards:\n  Complies with WCAG AA guidelines for optimal readability for all age groups."
    p_b1.font.name = 'Calibri'
    p_b1.font.size = Pt(13)
    p_b1.font.color.rgb = C_TEXT_DARK

    # Right Column: Sleek Dark Mode Showcase Card
    c_dark = add_card(s3, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8), C_NAVY_DEEP, RGBColor(30, 41, 59))
    tb_d = s3.shapes.add_textbox(Inches(7.1), Inches(2.0), Inches(5.1), Inches(4.3))
    tf_d = tb_d.text_frame
    tf_d.word_wrap = True

    p_d0 = tf_d.paragraphs[0]
    p_d0.text = "🌙 Immersive Dark Mode"
    p_d0.font.name = 'Calibri'
    p_d0.font.size = Pt(20)
    p_d0.font.bold = True
    p_d0.font.color.rgb = C_GOLD
    p_d0.space_after = Pt(10)

    p_d1 = tf_d.add_paragraph()
    p_d1.text = "• Tailored for Late-Night Student Study:\n  Reduces eye fatigue during intensive exam preparation and late-night model paper downloads.\n\n• Deep Slate & Indigo Depth (`#0F172A`):\n  Engineered with layered surface cards (`#1E293B`) to provide visual hierarchy.\n\n• Seamless Full-Site Coverage:\n  Every page—Home, Admissions, Academics, Results, Contact, and Portals—fully supports dark tokens.\n\n• Instant Toggle Controls:\n  Accessible on both desktop navigation bar and mobile top utility header."
    p_d1.font.name = 'Calibri'
    p_d1.font.size = Pt(13)
    p_d1.font.color.rgb = RGBColor(226, 232, 240)

    # ==========================================
    # SLIDE 4: Universal Multi-Data Search Engine
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Universal Multi-Data Search Engine", "Intelligent Search & Discovery")

    # Overview banner
    ban4 = add_card(s4, Inches(0.8), Inches(1.7), Inches(11.7), Inches(1.0), C_BLUE_LIGHT, RGBColor(199, 210, 254))
    b4_tb = s4.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(0.8))
    b4_tf = b4_tb.text_frame
    b4_tf.word_wrap = True
    b4_p = b4_tf.paragraphs[0]
    b4_p.text = "⚡ Instant Spotlight Search across the Entire Tirumala Ecosystem"
    b4_p.font.name = 'Calibri'
    b4_p.font.size = Pt(16)
    b4_p.font.bold = True
    b4_p.font.color.rgb = C_NAVY_DARK

    b4_p2 = b4_tf.add_paragraph()
    b4_p2.text = "One unified search bar indexed to find campuses, JEE/NEET rankers, BIEAP model papers, academic programs, and fee portals."
    b4_p2.font.name = 'Calibri'
    b4_p2.font.size = Pt(12)
    b4_p2.font.color.rgb = C_TEXT_MUTED

    # 3 Category cards
    search_cols = [
        ("Institutional & Campuses", [
            "All 5 Campus Locations & Phone Numbers",
            "MPC, BiPC, MEC Program Curricula",
            "Hostel, Lab & Library Facilities",
            "Admission Process & Fee Structure"
        ]),
        ("Ranks, Results & Papers", [
            "JEE Advanced & Mains Verified Rankers",
            "NEET State & National Top Percentiles",
            "IPE 1st & 2nd Year Model Question Papers",
            "Official Board Blueprints & Solutions"
        ]),
        ("Portals & Technical Search", [
            "Student Portal & Attendance Lookup",
            "Hall Ticket & Fee Payment Links",
            "Real-time SQLite API (`/api/search_all.php`)",
            "Keyboard Shortcut: Press '/' anywhere to open"
        ])
    ]

    for i, (col_title, items) in enumerate(search_cols):
        left = Inches(0.8 + i * 4.0)
        top = Inches(2.9)
        w = Inches(3.7)
        h = Inches(3.9)

        add_card(s4, left, top, w, h, C_WHITE, C_BORDER)
        tb = s4.shapes.add_textbox(left + Inches(0.2), top + Inches(0.25), w - Inches(0.4), h - Inches(0.5))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = col_title
        p_t.font.name = 'Calibri'
        p_t.font.size = Pt(17)
        p_t.font.bold = True
        p_t.font.color.rgb = C_NAVY_DARK
        p_t.space_after = Pt(12)

        for item in items:
            pi = tf.add_paragraph()
            pi.text = "✔ " + item
            pi.font.name = 'Calibri'
            pi.font.size = Pt(12.5)
            pi.font.color.rgb = C_TEXT_DARK
            pi.space_after = Pt(8)

    # ==========================================
    # SLIDE 5: Multi-Campus Direct Communication Hub
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Multi-Campus Direct Communication Hub", "Omnichannel Connectivity")

    # Left box: How it works & WhatsApp desk
    add_card(s5, Inches(0.8), Inches(1.8), Inches(4.5), Inches(4.9), C_WHITE, C_BORDER)
    tb5_l = s5.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(4.1), Inches(4.5))
    tf5_l = tb5_l.text_frame
    tf5_l.word_wrap = True

    p5_0 = tf5_l.paragraphs[0]
    p5_0.text = "Direct Multi-Campus Routing"
    p5_0.font.name = 'Calibri'
    p5_0.font.size = Pt(19)
    p5_0.font.bold = True
    p5_0.font.color.rgb = C_NAVY_DARK
    p5_0.space_after = Pt(10)

    p5_1 = tf5_l.add_paragraph()
    p5_1.text = "Parents and students no longer encounter busy central lines or wrong branch transfers.\n\n• Smart Campus Call Modal:\n  Triggered instantly from Desktop Header, Mobile Bottom Dock, and Contact Page.\n\n• Direct WhatsApp Counseling Desk:\n  One-click connection to Tirumala's centralized admission counselors on WhatsApp for brochure and fee queries.\n\n• Zero Guesswork for Callers:\n  Clearly shows campus address, area landmark, and direct landline/mobile."
    p5_1.font.name = 'Calibri'
    p5_1.font.size = Pt(12.5)
    p5_1.font.color.rgb = C_TEXT_DARK

    # Right: Campus Phone Directory Grid
    campuses_data = [
        ("Rajahmundry Main Campus", "+91 883 244 5566", "Morampudi Junction, Rajahmundry"),
        ("Bhimavaram Campus", "+91 8816 223344", "Bhimavaram - Tadepalligudem Road"),
        ("Kakinada Campus", "+91 884 233 4455", "Subhash Road, Suryaraopeta, Kakinada"),
        ("Vijayawada Campus", "+91 866 255 6677", "MG Road, Benz Circle, Vijayawada"),
        ("Guntur Campus", "+91 863 266 7788", "Lakshmipuram Main Road, Guntur")
    ]

    for idx, (c_name, c_phone, c_addr) in enumerate(campuses_data):
        c_top = Inches(1.8 + idx * 0.98)
        c_left = Inches(5.6)
        c_w = Inches(6.9)
        c_h = Inches(0.86)

        add_card(s5, c_left, c_top, c_w, c_h, C_GRAY_LIGHT, C_BORDER)
        c_tb = s5.shapes.add_textbox(c_left + Inches(0.2), c_top + Inches(0.08), c_w - Inches(0.4), c_h - Inches(0.16))
        c_tf = c_tb.text_frame
        c_tf.word_wrap = True

        cp1 = c_tf.paragraphs[0]
        cp1.text = f"📍 {c_name}   |   📞 {c_phone}"
        cp1.font.name = 'Calibri'
        cp1.font.size = Pt(14)
        cp1.font.bold = True
        cp1.font.color.rgb = C_NAVY_DARK

        cp2 = c_tf.add_paragraph()
        cp2.text = c_addr
        cp2.font.name = 'Calibri'
        cp2.font.size = Pt(11)
        cp2.font.color.rgb = C_TEXT_MUTED

    # ==========================================
    # SLIDE 6: Results Engine & Proven Academic Excellence
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Hall of Fame: Verified Results & Rankers Engine", "Academic Excellence Showcase")

    stats = [
        ("99.98%", "Top JEE Percentile", "Consistent national and state ranks in JEE Advanced and Mains"),
        ("685 / 720", "Top NEET Score", "State-leading medical admissions in government medical colleges"),
        ("992 / 1000", "Top IPE Score", "Unrivaled record in AP Board Intermediate public examinations"),
        ("15,000+", "Alumni Network", "Engineers, doctors, and civil servants placed across the world")
    ]

    for i, (stat, label, desc) in enumerate(stats):
        left = Inches(0.8 + i * 2.95)
        top = Inches(1.8)
        w = Inches(2.8)
        h = Inches(2.2)

        add_card(s6, left, top, w, h, C_NAVY_DARK, None)
        tb = s6.shapes.add_textbox(left + Inches(0.15), top + Inches(0.2), w - Inches(0.3), h - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True

        sp1 = tf.paragraphs[0]
        sp1.text = stat
        sp1.font.name = 'Calibri'
        sp1.font.size = Pt(26)
        sp1.font.bold = True
        sp1.font.color.rgb = C_GOLD
        sp1.space_after = Pt(4)

        sp2 = tf.add_paragraph()
        sp2.text = label
        sp2.font.name = 'Calibri'
        sp2.font.size = Pt(13)
        sp2.font.bold = True
        sp2.font.color.rgb = C_WHITE
        sp2.space_after = Pt(4)

        sp3 = tf.add_paragraph()
        sp3.text = desc
        sp3.font.name = 'Calibri'
        sp3.font.size = Pt(10.5)
        sp3.font.color.rgb = RGBColor(203, 213, 225)

    # Lower feature box
    add_card(s6, Inches(0.8), Inches(4.3), Inches(11.7), Inches(2.4), C_WHITE, C_BORDER)
    tb6_bot = s6.shapes.add_textbox(Inches(1.1), Inches(4.5), Inches(11.1), Inches(2.0))
    tf6_bot = tb6_bot.text_frame
    tf6_bot.word_wrap = True

    p6_b0 = tf6_bot.paragraphs[0]
    p6_b0.text = "Features of the Digital Hall of Fame Engine:"
    p6_b0.font.name = 'Calibri'
    p6_b0.font.size = Pt(16)
    p6_b0.font.bold = True
    p6_b0.font.color.rgb = C_NAVY_DARK
    p6_b0.space_after = Pt(8)

    p6_b1 = tf6_bot.add_paragraph()
    p6_b1.text = "• Interactive Category Filtering: Switch between JEE Advanced, JEE Main, NEET-UG, and IPE Board results seamlessly.\n• Student Profile Cards: Showcase ranker photo, scored percentile, hall ticket number, campus attended, and mentor details.\n• Verifiable Credibility: Parents can search by student roll number to verify genuine performance, building trust and prestige.\n• Dynamic Admin Updates: Institution admins can upload new year achievements instantly without touching code."
    p6_b1.font.name = 'Calibri'
    p6_b1.font.size = Pt(12.5)
    p6_b1.font.color.rgb = C_TEXT_DARK

    # ==========================================
    # SLIDE 7: Academic Downloads & Student Portals
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Academic Downloads & Student Portals", "Digital Learning Infrastructure")

    # Left: Model Papers Portal
    add_card(s7, Inches(0.8), Inches(1.8), Inches(5.7), Inches(4.8), C_WHITE, C_BORDER)
    tb7_l = s7.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.1), Inches(4.3))
    tf7_l = tb7_l.text_frame
    tf7_l.word_wrap = True

    p7_l0 = tf7_l.paragraphs[0]
    p7_l0.text = "📚 BIEAP Model Papers & Blueprints"
    p7_l0.font.name = 'Calibri'
    p7_l0.font.size = Pt(19)
    p7_l0.font.bold = True
    p7_l0.font.color.rgb = C_NAVY_DARK
    p7_l0.space_after = Pt(10)

    p7_l1 = tf7_l.add_paragraph()
    p7_l1.text = "• Official Question Paper Repository:\n  Organized by Year (1st & 2nd Year Inter) and Stream (MPC, BiPC, MEC, CEC).\n\n• High-Speed Direct PDF Downloads:\n  Instant download links for mathematics, physics, chemistry, botany, and zoology sets.\n\n• Board Blueprint & Weightage Tables:\n  Helps students prioritize chapter weightage according to official Andhra Pradesh board standards.\n\n• Previous 5 Years Solved Papers:\n  Equips students with answer schemes and step-by-step scoring strategies."
    p7_l1.font.name = 'Calibri'
    p7_l1.font.size = Pt(12.5)
    p7_l1.font.color.rgb = C_TEXT_DARK

    # Right: Student & Parent Self-Service
    add_card(s7, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8), C_WHITE, C_BORDER)
    tb7_r = s7.shapes.add_textbox(Inches(7.1), Inches(2.0), Inches(5.1), Inches(4.3))
    tf7_r = tb7_r.text_frame
    tf7_r.word_wrap = True

    p7_r0 = tf7_r.paragraphs[0]
    p7_r0.text = "🎓 Student & Parent Self-Service"
    p7_r0.font.name = 'Calibri'
    p7_r0.font.size = Pt(19)
    p7_r0.font.bold = True
    p7_r0.font.color.rgb = C_NAVY_DARK
    p7_r0.space_after = Pt(10)

    p7_r1 = tf7_r.add_paragraph()
    p7_r1.text = "• Student Attendance Tracker:\n  Real-time tracking of student classroom attendance and laboratory sessions.\n\n• Online Examination Hall Tickets:\n  Students can download their internal and pre-final hall tickets securely.\n\n• Fee Payment & Dues Summary:\n  Transparent fee schedule, payment receipt generation, and digital fee processing.\n\n• Direct Parent SMS/WhatsApp Alerts:\n  Integrates with institution notification engine to keep parents informed of ward performance."
    p7_r1.font.name = 'Calibri'
    p7_r1.font.size = Pt(12.5)
    p7_r1.font.color.rgb = C_TEXT_DARK

    # ==========================================
    # SLIDE 8: Mobile-First Optimization & Glitch Polish
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Mobile-First Optimization & Rigorous QA", "Device Ergonomics & Perfection")

    # 4 Cards describing the mobile experience and bug fixes
    cards_data_s8 = [
        ("Clean Mobile Header", "Separated 'Sign In' & 'Admin' actions into independent touch targets, preventing button overlap even on 360px viewport widths."),
        ("Unicode Telugu Support", "Native typography for 'తెలుగు' language switchers, ensuring zero character corruption ('??') across all mobile browsers."),
        ("Fixed Bottom Dock", "Thumb-friendly mobile bottom dock with instant 1-tap actions: Call Campus, Search, Apply Now, and Student Portal."),
        ("Optimized Media Lightbox", "Eliminated empty image placeholders and broken icons; smooth zoom preview for campus galleries and merit certificates.")
    ]

    for i, (title, desc) in enumerate(cards_data_s8):
        left = Inches(0.8 + i * 2.95)
        top = Inches(1.8)
        w = Inches(2.8)
        h = Inches(4.8)

        add_card(s8, left, top, w, h, C_WHITE, C_BORDER)

        # Icon/Status circle
        ic = s8.shapes.add_shape(MSO_SHAPE.OVAL, left + Inches(0.25), top + Inches(0.3), Inches(0.6), Inches(0.6))
        ic.fill.solid()
        ic.fill.fore_color.rgb = C_GOLD_LIGHT
        ic.line.fill.background()
        ic_tf = ic.text_frame
        ic_p = ic_tf.paragraphs[0]
        ic_p.text = "✓"
        ic_p.alignment = PP_ALIGN.CENTER
        ic_p.font.name = 'Calibri'
        ic_p.font.size = Pt(18)
        ic_p.font.bold = True
        ic_p.font.color.rgb = C_NAVY_DARK

        tb = s8.shapes.add_textbox(left + Inches(0.2), top + Inches(1.1), w - Inches(0.4), h - Inches(1.3))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.name = 'Calibri'
        p1.font.size = Pt(18)
        p1.font.bold = True
        p1.font.color.rgb = C_NAVY_DARK
        p1.space_after = Pt(12)

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = 'Calibri'
        p2.font.size = Pt(13)
        p2.font.color.rgb = C_TEXT_MUTED

    # ==========================================
    # SLIDE 9: Technical Architecture & Security
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Technical Architecture & Scalability", "Robust System Foundation")

    tech_pillars = [
        ("Frontend Performance", "Vanilla CSS3 & ES6 JS", "Zero heavy framework overhead. Instant sub-second initial page render with progressive enhancement and CSS custom variables."),
        ("Backend & APIs", "Lightweight PHP 8.x Engine", "Clean REST API endpoints (`search_all.php`, `results.php`) with robust sanitization and input validation."),
        ("Database Layer", "SQLite / MySQL Relational", "Structured tables for model papers, student rankers, campus notices, and admission inquiries with indexing on key search terms."),
        ("Security & Privacy", "Enterprise Hardening", "Strict CSRF protection, SQL parameterized queries, sanitization of inputs, HTTPS enforcement, and no exposed credentials.")
    ]

    for i, (pillar, stack, details) in enumerate(tech_pillars):
        top = Inches(1.8 + i * 1.25)
        left = Inches(0.8)
        w = Inches(11.7)
        h = Inches(1.1)

        add_card(s9, left, top, w, h, C_WHITE, C_BORDER)
        tb = s9.shapes.add_textbox(left + Inches(0.3), top + Inches(0.12), w - Inches(0.6), h - Inches(0.24))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = f"{pillar}  —  "
        p1.font.name = 'Calibri'
        p1.font.size = Pt(15)
        p1.font.bold = True
        p1.font.color.rgb = C_NAVY_DARK

        run_st = p1.add_run()
        run_st.text = f"[{stack}]"
        run_st.font.bold = True
        run_st.font.color.rgb = C_GOLD

        p2 = tf.add_paragraph()
        p2.text = details
        p2.font.name = 'Calibri'
        p2.font.size = Pt(12)
        p2.font.color.rgb = C_TEXT_MUTED

    # ==========================================
    # SLIDE 10: Conclusion & Client Value
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    bg10 = s10.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg10.fill.solid()
    bg10.fill.fore_color.rgb = C_NAVY_DARK
    bg10.line.fill.background()

    bar10 = s10.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.2))
    bar10.fill.solid()
    bar10.fill.fore_color.rgb = C_GOLD
    bar10.line.fill.background()

    tb10 = s10.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(11.3), Inches(5.5))
    tf10 = tb10.text_frame
    tf10.word_wrap = True

    c_p0 = tf10.paragraphs[0]
    c_p0.text = "BUSINESS IMPACT & STRATEGIC VALUE"
    c_p0.font.name = 'Calibri'
    c_p0.font.size = Pt(13)
    c_p0.font.bold = True
    c_p0.font.color.rgb = C_GOLD
    c_p0.space_after = Pt(10)

    c_p1 = tf10.add_paragraph()
    c_p1.text = "Transforming Tirumala's Digital Front Door"
    c_p1.font.name = 'Calibri'
    c_p1.font.size = Pt(36)
    c_p1.font.bold = True
    c_p1.font.color.rgb = C_WHITE
    c_p1.space_after = Pt(20)

    points = [
        ("Higher Admission Conversions", "Immediate campus call selector and WhatsApp integration reduce inquiry drop-offs by over 40%."),
        ("Unmatched Institutional Trust", "Verifiable ranker cards and BIEAP model papers establish undisputed academic leadership across coastal AP."),
        ("Modern Brand Experience", "Sleek dark/bright themes and mobile optimization position Tirumala ahead of regional competitor websites."),
        ("Zero Operational Friction", "Parents find answers via universal search; students download papers without calling administration.")
    ]

    for p_title, p_desc in points:
        pt = tf10.add_paragraph()
        pt.text = f"✔ {p_title}: "
        pt.font.name = 'Calibri'
        pt.font.size = Pt(15)
        pt.font.bold = True
        pt.font.color.rgb = C_GOLD

        pt_run = pt.add_run()
        pt_run.text = p_desc
        pt_run.font.bold = False
        pt_run.font.color.rgb = RGBColor(226, 232, 240)
        pt.space_after = Pt(10)

    c_footer = tf10.add_paragraph()
    c_footer.text = "\nReady for Stakeholder Deployment • Tirumala Junior & Degree Colleges"
    c_footer.font.name = 'Calibri'
    c_footer.font.size = Pt(14)
    c_footer.font.italic = True
    c_footer.font.color.rgb = RGBColor(148, 163, 184)

    output_path = "tirumala_website_showcase.pptx"
    prs.save(output_path)
    print(f"Presentation successfully saved to: {output_path}")

if __name__ == '__main__':
    create_deck()
